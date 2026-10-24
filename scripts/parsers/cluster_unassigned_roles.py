#!/usr/bin/env python3
"""
cluster_unassigned_roles.py
----------------------------
Takes the persons that parse_person_role.py leaves WITHOUT any role bucket
(person_role.csv, empty `roller`) and groups them by the profession words
in their description into at most 20 clusters, so a human annotator can
assign one of the existing buckets (or "ingen passende bucket") to each
cluster instead of reviewing ~4,700 persons one by one.

The clusters are deliberately NOT the 9 buckets. They are finer-grained
(and include groups no bucket covers, e.g. servants, travellers, crafts)
so that each cluster can be mapped to a single bucket — or flagged as
having none — without splitting it.

Same referent-safety as parse_person_role.py: the description is run
through RELATION_RE / MARRIED_RE first, so "g. 1852 m. Pianisten Otto
Goldschmidt" on Jenny Lind's entry does NOT put her in a pianist cluster.

Reads : data/normalized/entities.csv, data/normalized/person_role.csv
Writes: data/normalized/person_role_unassigned_clusters.csv  (one row/cluster)
        data/normalized/person_role_unassigned_persons.csv   (one row/person)
        data/normalized/person_role_unassigned_annotation.xlsx
                                  (same two tables; `bucket` column with
                                   a drop-down of the existing buckets)

A person gets one PRIMARY cluster (the profession word that comes first in
the cleaned description) plus the list of other clusters that also matched.
Persons with no profession word are reported as residual groups, not as
clusters.
"""

import csv
import os
import re
import sys
from collections import Counter, defaultdict

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, os.path.dirname(__file__))
import parse_person_role as ppr  # noqa: E402  (reuse RELATION_RE/MARRIED_RE)
import _person_groups  # noqa: E402  (scripts/build_mockup is put on sys.path by ppr)

ENTITIES = os.path.join(ROOT, "data", "normalized", "entities.csv")
ROLES = os.path.join(ROOT, "data", "normalized", "person_role.csv")
OUT_CLUSTERS = os.path.join(ROOT, "data", "normalized", "person_role_unassigned_clusters.csv")
OUT_PERSONS = os.path.join(ROOT, "data", "normalized", "person_role_unassigned_persons.csv")
OUT_XLSX = os.path.join(ROOT, "data", "normalized", "person_role_unassigned_annotation.xlsx")

NO_BUCKET = "Ingen passende bucket"

# Cluster id, name, token regex. A token is a lower-cased word of the
# cleaned description (hyphenated words are also tried part by part);
# patterns are anchored on the token and written as head-noun suffixes
# where Danish compounding makes that safe (…maler, …fabrikant, …kirurg).
# First listed cluster wins a single token.
CLUSTERS = [
    ("K01", "Musik, sang og scene",
     r".*sanger(inde)?|sangpædagog|sangmester|.*sangerinde|musiker(inde)?|(violin|klarinet|cello|violoncel|fløjt|piano|harpe|horn|trompet|pauk|kontrabas|fagot|oboe|guitar|mandolin|"
     r"koncert)ist(inde)?|klarinettist|pianist(inde)?|violinist(inde)?|cellist|fløjtenist|harpenist|guitarist|"
     r".*virtuos(er)?|.*spiller(inde|ske)?|koncertmester|.*kapelmester|.*komponist|musiklærer(inde)?|"
     r"musikdirektør|brigade-musikdirektør|.*skuespiller(inde)?|.*danser(inde)?|tryllekunstner|komiker|"
     r"balletmester|statist|teaterdirektør|teaterentreprenør|instruktør|kunstrytter|artist|korist(inde)?|"
     r"ornatrice|figurant|hornblæser|improvisatrice|recitatrice|magiker|hukommelseskunstner|"
     r"kunstskøjteløber|skøjteløberdanser|sortekunstner|.*skuespillerinde|organist|domorganist|kvartet|"
     r"damekvartet|vidunderbørn|musikforlægger|musikanmelder|trommevirtuos|hofteaterintendant"),
    ("K02", "Maleri, tegning, skulptur, foto og arkitektur",
     r".*maler(inde)?|.*malerinde|billedhugger(inde)?|billedskærer|.*tegner(inde)?|kobberstikker|xylograf|"
     r"litograf|kunstner(inde)?|graver|overgraver|.*fotograf|fotografer|arkitekt|tegnelærer|silhouetklipper|"
     r"kunstgibser|stukkatør|træskærer|kartograf|karikaturtegner|kunstkritiker"),
    ("K03", "Digtning, forfatterskab og presse",
     r"digter(inde)?|poet|forfatterverbum|.*forfatter(inde)?|forf|oversætter(inde)?|translatør|skribent|redaktør|"
     r"journalist|.*digter|dramatiker|romanskriver|libretist|medudgiver|tidsskriftsudgiver|litteraturanmelder|"
     r"rigsdagsstenograf|bibliograf|redaktionssekretær|.*historiker|kulturhistoriker|krigshistoriker|"
     r"litteraturhistoriker|kirkehistoriker|musikhistoriker|korrektør|boghandlerfuldmægtig(?!x)"),
    ("K04", "Videnskab og lærdom",
     r"professor|docent|universitetsdocent|kemiker|astronom|geolog|orientalist|konkyliolog|botaniker|zoolog|"
     r"filolog|arkivar|bibliotekar|underbibliotekar|overbibliotekar|amanuensis|naturforsker|opdagelsesrejsende|"
     r"etnograf|arkæolog|antikvar|sprogforsker|statistiker|mineralog|entomolog|ornitolog|fysiker|matematiker|"
     r"lærde|videnskabsmand|geograf|numismatiker|forsker|filosof|antropolog|genealog|hippolog|oldgransker|"
     r"inspektor|fysikus|polytekniker|manuduktør|analytiker|galvanoplastiker|opdager|krigsforsker"),
    ("K05", "Medicin, apotek og teknik",
     r".*kirurg|tandlæge|tandmedikus|dyrlæge|apoteker|jordemoder|læge|lægen|.*læge|.*overlæge|ingeniør|"
     r".*ingeniør|civilingeniør|opfinder|erfinder|maskinmester|landinspektør|landmaaler|konstruktør|"
     r"barselskone|sygeplejerske|chirurg|fysikus|stadsfysikus|landfysikus|medicus|vaagekone|badekone|"
     r"orgelbygger|skibsbygger|mekaniker|fyringeniør|maskinist|patient"),
    ("K06", "Undervisning og opdragelse",
     r".*lærer(inde)?|.*lehrer|informator|guvernante|opdrager(inde)?|skolebestyrer(inde)?|institutbestyrerinde|"
     r"forstanderinde|forstander|.*forstander|skoleinspektør|rektor|pædagog|opdragelsesanstalt|"
     r"skolemester|skoleholder|skolemand|kateket|seminarist|lærerinden"),
    ("K07", "Akademisk grad, studium og elevtid",
     r"cand|phil|theol|jur|polyt|stud|exam|candidat|studerende|student|universitet|doktor|licentiat|elev|.*elev"),
    ("K08", "Kirke og gejstlighed",
     r"ærkebiskop|kardinal|munk|diakonisse|klokker|apostel|helgen|helgeninde|reformator|kirkereformator|nonne|"
     r"abbed|abbedisse|missionær|degn|kirkebetjent|.*præst|præsten|fyrst-abbed|pastor|rabbiner|kirketjener|pave|patriark|"
     r"eremit|martyr|eneboer|pater|priorinde|superintendent|ypperstepræst|prior|storpriorats-sekretær"),
    ("K09", "Adel, fyrstehuse og hof",
     r"adelsmand|adelsdame|adelig|adelige|stormand|herremand|fyrste|fyrstinde|kurfyrste|kurfyrstinde|landgreve|"
     r"kejser|kejserinde|hertug|hertugen|hertuginde|erkehertuginde|prins|prinser|prinsesse|baronesse|baron|"
     r"komtesse|friherre|friherreinde|overhofmester|staldmester|overstaldmester|hofdame|hofmarskal|jarl|"
     r"marquess|marquis|lady|emir|pasha|kong|konge|dronning|greve|grevinde|rigsgrevinde|grev|ridder|"
     r"lensgreve|kammerherre|hofjægermester|overjægermester|jægermester|regent|regentinde|tronarving|"
     r"tronprætendent|storhertug|storhertuginde|sultan|shah|infantinde|rigshofmester|rigsmarsk|oberstkämmerer|"
     r"kämmerer|majoratsbesidder|fideikommisbesidder|kabinetskammerherre|væbner|jagtjunker|kammerjunker|"
     r"hofkapelmester(?!x)|mulatinde(?!x)"),
    ("K10", "Kloster og stift",
     r"kloster|klosterfrøken|klosterfrøknerne|jomfrukloster|dekanesse|klosterintendant|klosterkasserer|"
     r"abbedissen|stiftsdame|konventualinde|vemmetofte|vallø|klosterprovst|stift"),
    ("K11", "Militær og krig",
     r"major|menig|kompagni|infanteri-regiment|.*regiment|feltmarskal|kadet|søkadet|hærfører|.*kaptajn|"
     r"hauptmann|frihedskæmper|soldat|artilleri|kavaleri|marskal|officer|korporal|sergent|sergeant|"
     r"fændrik|fanejunker|rytter|kommandør|.*løjtnant|oberst|general|admiral|.*kommissær|livjæger|husar|"
     r"dragon|adjudant|generaladjudant|veteran|militær|oberstløjtnant|søofficer|kornet|skarpskytte|underofficer|"
     r"kommandant|batteri|bataillon|armeen|hæren|artillerichef|høvedsmand|rustmester|felttøjmester|"
     r"feltkasserer|krieger|gendarmeriofficer|regimentskvartermester|reserveofficersaspirant|krigsminister|"
     r"frikorpsfører|bergmanns-lieutenant|kolonel-intendant|overauditør|flaadefører|marineminister"),
    ("K12", "Embedsværk, retsvæsen og politik",
     r".*raad|rigsraad|statholder|senator|geheimerat|geheimer|.*assessor|.*advokat|advokat|.*dommer|dommer|"
     r"notar|sorenskriver|.*skriver|politimester|statsminister|ministerpræsident|.*minister|minister|"
     r".*kansler|chargé|konsul|gesandt|ambassadør|toldkontrollør|.*toldinspektør|.*told.*|postmester|"
     r"postvæsenet|lensmand|amtmand|borgmester|overborgmester|byfoged|herredsfoged|landfoged|foged|"
     r"politibetjent|embedsmand|.*embedsmand|generaldirektoratet|præsident|president|präsident|politiker|"
     r".*politiker|folketingsmand|stænderdeputeret|revisor|kancellist|rådhuskanslist|departementschef|"
     r"stiftamtmand|landsdommer|prokurator|justitiarius|kontrollør|.*kontrollør|overkontrollør|bureauchef|"
     r"sekretær|.*sekretær|.*sekretær|politi|diplomat|legationssekretær|.*inspektør|.*postinspektør|"
     r"postdirektør|postekspeditør|postexpedient|post-officiant|jembanedirektør|departementsdirektør|"
     r"generaldirektør|revisionschef|rentemester|skatmester|kronsyndikus|civilguvernør|stadssekretær|"
     r"statssekretær|regeringspræsident|retspræsident|lovgiver|trafikassistent|havnemester|statistformand|"
     r"stempelforvalter|kortstemplingsforvalter|konsumptionskasserer|toldkasserer|toldbetjent|"
     r"landskamrer|landfysikus|oprørsleder|gehejmekabinetsminister|gehejmestatsminister|justitsminister|"
     r"finansminister|sagfører|ministerialbud|politiinspektør|borger|socialist|volontør|næringssorger|kæmner"),
    ("K13", "Handel, bank og industri",
     r"bankier|bankierfirma|firma|.*købmand|.*kobmand|.*handler|.*fabrikant|fabrikejer|fabriksejer|.*fabrik|"
     r"vekselerer|veksellerer|vekselmægler|finansmand|forretningsmand|medindehaver|medejer|medbesidder|"
     r"medbesidderinde|medbestyrerinde|handelshus|assurandør|assurandører|assurancemægler|kasserer|"
     r"billetkasserer|kontorist|bogholder|.*bogholder|bogtrykker|bogtrykkeri|forlagsboghandel|boghandler|"
     r"direktør|direktor|director|meddirektør|lotteridirektør|grosserer|vingrosserer|entreprenør|"
     r"bade-entreprenør|rederi|skibsreder|reder|mægler|.*mægler|agent|handelsagent|vin-agent|kommissionær|"
     r"fuldmægtig|handelsfuldmægtig|handelsmand|bankdirektør|bankchef|forlag|bogforlag|storkøbmand|"
     r"kommis|handelsrejsende|isenkræmmer|hørkræmmer|urtekræmmer|kræmmer|"
     r"klasselotterikollektør|lotterikollektrice|speditør|manager|afdelingsbestyrer|forretningsfører|"
     r"regnskabsfører|jernstøberiejer|mølleejer|mineejer|brygger|galanteriehändler|sukkerraffinadør|"
     r"sukkerhusforvalter|pakhusforvalter|blanketforvalter|handelsgartner|handelskontorist|brukspatron|"
     r"ekspeditrice|indehaverske|.*medhjælper|værkmester|fabrikmester|ordensjuvelerer|hofjuvelerer|juvelerer|hofparfumør"),
    ("K14", "Håndværk og laugsfag",
     r".*mager|.*magermester|smed|.*smed|.*smedemester|skrædermester|skræddermester|skrædder|skræder|"
     r"skræddersvend|guldsmed|konditor|frisør|frisørmester|hoffrisør|barber|typograf|tømrer|murer|snedker|"
     r".*snedkersvend|bager|hofbagermester|slagter|brændevinsbrænder|malermester|.*svend|bødker|garver|"
     r"farver|væver|bogbinder|drejer|tapetserer|glarmester|murermester|tømrermester|snedkermester|"
     r"blikkenslager|reberbaner|sejlmager|gørtler|rebslager|sæbesyder|pottemager|kunstdrejer|"
     r"skomagermester|skomagersvend|hof-skrædermester|plattenslager|jernstøber|dampmøller|"
     r"kældermester|overbanemester"),
    ("K15", "Værtshus, hotel og pension",
     r"værtinde|vært|værtshusholder|restauratør|gæstgiver|gæstgiverske|hotelværtinde|pensionatsværtinde|"
     r"hotelejer|hotelejerske|hotelvært|hotelværter|værtshus|hotelier|kroejer|kromand|pensionatsholder|"
     r"hotelbestyrer|pensionatsbestyrerinde|café|caféejer|kaffehus|kelner|overtjener|"
     r"traktør|husejerske|hovmester"),
    ("K16", "Tjenestefolk og lavt lønnet arbejde",
     r"tjenestepige|tjenestedreng|tjener|kammertjener|kahytstjener|kusk|portner|opvarter|kokkepige|kok|stuepige|"
     r"hotelstuepige|husbestyrerinde|barnepige|amme|bondedreng|bydreng|fiskerdreng|jydepige|tigger|tiggerske|daglejer|arbejder|"
     r"lakaj|betjent|vægter|husjomfru|kammerpige|kammerjomfru|stalddreng|staldkarl|vaskekone|vaskerkone|"
     r"dagarbejder|bybud|drager|tjenestekarl|tyende|jungmand|sypige|syerske|husholderske|hønsekone|"
     r"oppasser|ledsager|kurer|kontorbud|bybud|opsynsmand|bjergmand|bjergværksarbejder|vognmand|aftjent|portnerkone|"
     r"bagger|beder"),
    ("K17", "Landbrug, skov og gods",
     r"gaardejer|gårdejer|godsejer|landmand|skovrider|skovfoged|overførster|oberförster|forstmand|forstmester|"
     r"over-forstmester|slotsgartner|gartner|overgartner|undergartner|gartnerelev|fisker|husfoged|proprietær|"
     r"husmand|gaardmand|bonde|jordbesidder|forpagter|avlsbruger|landbrug|forvalter|ladefoged|godsforvalter|"
     r"godsinspektør|forstinspektør|økonomiinspektør|hyrde|mejerist|landbrugskandidat|ejendomsbesidder|"
     r"heste-ejer|slotsforvalter|bygningsforvalter|jæger|bjørnejæger|under-kornmaaler|lensbogholder"),
    ("K18", "Søfart, jernbane, telegraf og trafik",
     r"skibsfører|fører|baadfører|sømand|matros|lods|styrmand|skipper|konduktør|diligence-konduktør|postfører|"
     r"dampskibsfører|havnefoged|havnekontrollør|fyrvogter|postillon|dampskibskaptajn|orlogsmand|søfarende|"
     r"rorsmand|telegrafist|telegrafbestyrer|skibsklarerer|luftskipper|jernbaneingeniør|vandbygningsingeniør|vandbygningsdirektør"),
    ("K19", "Stiftere, mæcener, filantroper og samlere",
     r"filantrop|mæcen|kunstmæcen|stifter|grundlægger|medstifter|legatstifter|legatstifterske|samler|"
     r"autografsamler|malerisamler|godssamlerske|protektor|velynder|velgører|donator|forkæmper|"
     r"grundlæggeren|reformator"),
    ("K20", "Ikke-erhverv: rejsende, gæster og privatpersoner",
     r"rejsende|passager|diligencepassager|jernbanepassager|jembanepassager|diligence-passager|tog-passager|"
     r"dampskibspassager|turist|turister|hotelgæst|gæst|gæster|badegæster|logerende|emigrant|udvandrer|"
     r"pensionær|pensionærer|partikulier|rentier"),
]

# Clusters that only become the primary one when nothing else matched.
LOW_PRIORITY = {"K07", "K20"}

# Weak, workplace-only words: used ONLY when a person has no strong
# profession word, because "ved Det kgl. Teater" says where, not what.
WEAK = {
    "teater": "K01", "hofteatret": "K01", "hoftheater": "K01", "kapel": "K01", "casino": "K01",
    "museum": "K04", "bibliotek": "K04", "hospital": "K05", "skole": "K06", "seminarium": "K06",
}

# Words that describe the person's place in a family, not a profession.
RELATION_ONLY_RE = re.compile(
    r"elskerinde|veninde|beundrerinde|beundrer|slægtning|tante|svoger|svigersøn|svigermoder|svigerinde|"
    r"svigerdatter|svigerfader|kusine|fætter|nevø|niece|niecer|plejedatter|moder|fader|datter|søn|sønner|"
    r"broder|søster|søstre|enke|enkefrue|frue|frøken|frøknerne|frøkner|døtre|familie|familien|tvilling|"
    r"ægteskabet|forlovet|separeret|dattersøn|halvbroder|halvsøster|farbroder|morbroder|moster|fatter|"
    r"præstekone|bondekone|portnerkone|gaardmandskone|præstedatter|borgerdatter|broderdatter|"
    r"søsterdatter|.*datter|generalinde|hofraadinde|kammerjunkerinde|pige|dreng|piger"
)
# Generic leader/manager/assistant words whose domain depends on the
# (missing) workplace — reported as a residual group, not a cluster.
GENERIC_RE = re.compile(r"leder|chef|bestyrer|bestyrerinde|formand|formynder|assistent|konsulent|medhjælper|deltager")

# Names/places/months etc. are never professions, even if a loose regex
# above (e.g. `.*raad`, `.*told.*`) would let them through.
STOP = {
    "senere", "ikke", "eller", "enten", "hans", "efter", "august", "juli", "april", "juni", "marts",
    "september", "november", "december", "oktober", "januar", "februar", "maj", "ansat", "under", "mellem",
    "arnim", "fornem", "slags", "synes", "første", "hvis", "hvem", "hvad", "være", "have", "havde",
    "skal", "mener", "holdt", "fundet", "nævnt", "angivet", "boede", "lever", "over", "derefter",
    "afsked", "afgøres", "misforstået", "måske", "nærmere", "foregav", "forskrevet", "opløst",
    "ophævet", "kaldes", "kaldet", "gift", "sidste", "store", "lille", "unge", "ældste", "yngste",
    "tidligere", "arvede", "jura", "med", "tønder", "mayer", "meyer", "müller", "richter", "ritter",
    "winter", "schiller", "walter", "fischer", "peter", "peder", "krüger", "köster", "töpler",
    "zöllner", "burmeister", "heiter", "hedder", "heller", "tanger", "banner", "baller", "bremer",
    "brenner", "behager", "bearbejder", "befrier", "bergmanns", "blätter", "bonnier", "boder",
    "colombier", "cowper", "devrient", "donner", "edler", "eroder", "fenger", "flytter", "gener",
    "haffner", "hilmer", "jasper", "jesper", "koller", "krogh", "mariager", "münster", "nørager",
    "ohlenschlager", "ohmeyer", "plätzer", "reimer", "richter", "ræder", "skirner", "spencer",
    "thelen", "tidemand", "trier", "wasser", "winther", "zøylner", "ønsker", "søger", "sværmer",
    "morder", "morderske", "hohenstaufer", "hengist", "hermoder", "birger", "holger", "karakter",
    "oplæsninger", "underholdninger", "sangforeninger", "værker", "forkæmper", "ledsager", "hannover",
    "hannover", "skelskør", "kongsvinger", "helsingør", "hallager", "broager", "grant", "beder",
    "bagger", "kurer", "drager", "reder", "manager", "kontor", "firma", "forlag", "møller",
}


def _compile():
    out = []
    for cid, name, pat in CLUSTERS:
        out.append((cid, name, re.compile(r"^(?:%s)$" % pat)))
    return out


TOKEN_RE = re.compile(r"[A-Za-zÆØÅæøåéüöäÉ][A-Za-zÆØÅæøåéüöäÉ'\-]*")


# Stricter than parse_person_role.py's cleaning, because here a leaked
# title would put a person into the wrong cluster:
#   * "g. 1° 1810 m. X, 2° 1822 m. Y" (numbered marriages) is stripped
#   * a clause like "Grev Julius Holcks Svigersøn" (genitive + kinship)
#     describes someone else, so the whole clause goes
#   * words after hos/af/til/for/under/efter up to the next comma are the
#     OBJECT of the role ("Kammertjener hos Kong Christian VIII",
#     "Stifter af Vemmetofte Kloster"), never the person's own role
#   * "Pige/Dreng hos|i X" is a servant; bare Pige/Dreng is just age/sex
NUMBERED_MARRIAGE_RE = re.compile(r"(?:\bg\.\s*)?\b\d\s*°\s*(?:\d{4}\s*)?m\.\s*[^,;]*[,;]?\s*", re.I)
GENITIVE_KIN_RE = re.compile(
    r"\b\w+(?:s|’s|'s)\s+(?:Svigersøn|Svigerdatter|Søn|Datter|Broder|Søster|Enke|Enkemand|Hustru|Moder|Fader|"
    r"Niece|Nevø|Plejedatter|Plejesøn|Kusine|Fætter|Svoger|Svigermoder|Svigerfader|Barnebarn|Dattersøn|"
    r"Sønnesøn|Mand)\b", re.I)
OBJECT_PREPS = {"hos", "af", "til", "for", "under", "efter"}
SERVANT_RE = re.compile(r"\b(Pige|Dreng)\s+(hos|i)\b")
VERB_WRITES_RE = re.compile(r"\bskriver\s+(Vers|Eventyr|Digte|Romaner)", re.I)


def clean(desc):
    d = ppr.MARRIED_RE.sub("", ppr.RELATION_RE.sub("", desc or ""))
    d = NUMBERED_MARRIAGE_RE.sub("", d)
    d = SERVANT_RE.sub(lambda m: "Tjeneste" + m.group(1).lower() + " " + m.group(2), d)
    d = VERB_WRITES_RE.sub("forfatterverbum ", d)
    clauses = re.split(r"([,;])", d)
    kept = [c for c in clauses if not GENITIVE_KIN_RE.search(c)]
    return "".join(kept)


def tokens(desc):
    out = []
    base = 0
    for clause in re.split(r"(?<=[,;])", clean(desc)):
        skip = False
        for m in TOKEN_RE.finditer(clause):
            t = m.group().lower().rstrip("-")
            if skip:
                continue
            out.append((base + m.start(), t))
            if t in OBJECT_PREPS:
                skip = True
        base += len(clause)
    return out


def _match(t, compiled):
    cands = [t] + [p for p in t.split("-") if len(p) >= 3] if "-" in t else [t]
    if t.endswith("-frue"):  # "Bankier-Frue": a banker's wife, not a banker
        return None
    for c in cands:
        if c in STOP or len(c) < 3:
            continue
        for cid, _name, rx in compiled:
            if rx.match(c):
                return cid
    return None


def classify(desc, compiled):
    """Returns (strong_hits, weak_hits, is_relation, is_generic) where
    hits are [(pos, cid, token)]."""
    toks = tokens(desc)
    strong, weak = [], []
    for pos, t in toks:
        cid = _match(t, compiled)
        if cid:
            strong.append((pos, cid, t))
        elif t in WEAK:
            weak.append((pos, WEAK[t], t))
    rel = any(RELATION_ONLY_RE.fullmatch(t) for _p, t in toks)
    gen = any(GENERIC_RE.fullmatch(t) for _p, t in toks)
    return strong, weak, rel, gen


def main():
    ents = {}
    with open(ENTITIES, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["entity_type"] == "person":
                ents[r["entity_id"]] = r
    roles = {}
    with open(ROLES, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            roles[r["entity_id"]] = r

    # Grupper (kønsgennemsynets »Andet«) får ingen rolle og hører ikke til i
    # den manuelle kø — build_persons_extra.py sætter dem på Endnu ukategoriseret.
    groups = _person_groups.load_group_ids()
    un = [e for eid, e in ents.items() if not roles[eid]["roller"] and eid not in groups]
    print(f"Personer i registret: {len(ents):,}")
    print(f"Uden rolle-bucket:    {len(un):,} (uden {len(groups & ents.keys()):,} grupper)")

    compiled = _compile()
    names = {cid: name for cid, name, _ in CLUSTERS}
    prim = Counter()
    anyc = Counter()
    term_counts = defaultdict(Counter)
    examples = defaultdict(list)
    persons_out = []
    residual = Counter()
    residual_ex = defaultdict(list)

    for e in un:
        desc = (e["description"] or "").strip()
        strong, weak, rel, gen = classify(desc, compiled)
        hits = strong if strong else weak
        cids_in_order = []
        for _pos, cid, _t in sorted(hits):
            if cid not in cids_in_order:
                cids_in_order.append(cid)
        # A degree ("dr.jur. 1870, Dialektdigter") or a travel label never
        # outranks a real occupation, wherever it stands in the text.
        cids_in_order.sort(key=lambda c: c in LOW_PRIORITY)  # stable
        if cids_in_order:
            p = cids_in_order[0]
            prim[p] += 1
            for c in cids_in_order:
                anyc[c] += 1
            for _pos, cid, t in hits:
                term_counts[cid][t] += 1
            if len(examples[p]) < 6:
                examples[p].append(e["label"])
            grp = ""
        else:
            p = ""
            if not desc:
                grp = "R1 Ingen beskrivelse (kun henvisning)"
            elif gen:
                grp = "R2 Generisk stilling uden domæne (leder/chef/assistent…)"
            elif rel:
                grp = "R3 Kun slægtskab/ægteskab/relation"
            elif not [t for _p, t in tokens(desc) if len(t) > 3]:
                grp = "R3 Kun slægtskab/ægteskab/relation"
            else:
                grp = "R4 Anden tekst uden erhvervsord (sted/dato/navn)"
            residual[grp] += 1
            if len(residual_ex[grp]) < 8:
                residual_ex[grp].append(f"{e['label']} | {desc[:80]}")
        persons_out.append({
            "entity_id": e["entity_id"],
            "label": e["label"],
            "description": desc,
            "klynge": p,
            "klynge_navn": names.get(p, ""),
            "andre_klynger": ";".join(cids_in_order[1:]),
            "matchede_termer": ";".join(
                ("~" if not strong else "") + t for _pos, _c, t in sorted(hits)),
            "restgruppe": grp,
        })

    with open(OUT_PERSONS, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(persons_out[0].keys()))
        w.writeheader()
        w.writerows(persons_out)

    cl_rows = []
    for cid, name, _ in CLUSTERS:
        top = ", ".join(f"{t} ({n})" for t, n in term_counts[cid].most_common(12))
        cl_rows.append({
            "klynge": cid,
            "navn": name,
            "antal_personer": prim[cid],
            "antal_inkl_sekundaere": anyc[cid],
            "top_termer": top,
            "eksempler": " ; ".join(examples[cid]),
            "bucket": "",
            "kommentar": "",
        })
    with open(OUT_CLUSTERS, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(cl_rows[0].keys()))
        w.writeheader()
        w.writerows(cl_rows)

    clustered = sum(prim.values())
    print(f"\nMed mindst ét erhvervsord (klyngbare): {clustered:,}")
    print("\nKlynger (primær tilhørsforhold):")
    for r in sorted(cl_rows, key=lambda r: -r["antal_personer"]):
        print(f"  {r['antal_personer']:5,}  {r['klynge']} {r['navn']}   [inkl. sekundære: {r['antal_inkl_sekundaere']}]")
    print("\nRest (ingen klynge):")
    for g, n in sorted(residual.items()):
        print(f"  {n:5,}  {g}")
        for x in residual_ex[g][:5]:
            print(f"           {x}")
    write_xlsx(cl_rows, persons_out)


def write_xlsx(cl_rows, persons_out):
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl ikke installeret — springer xlsx over")
        return
    buckets = list(ppr.BUCKET_ORDER) + [NO_BUCKET]
    wb = Workbook()
    ws = wb.active
    ws.title = "Klynger"
    cols = list(cl_rows[0].keys())
    ws.append(cols)
    for r in cl_rows:
        ws.append([r[c] for c in cols])
    lst = wb.create_sheet("Buckets")
    for b in buckets:
        lst.append([b])
    dv = DataValidation(type="list", formula1=f"=Buckets!$A$1:$A${len(buckets)}", allow_blank=True)
    ws.add_data_validation(dv)
    bcol = cols.index("bucket") + 1
    from openpyxl.utils import get_column_letter
    dv.add(f"{get_column_letter(bcol)}2:{get_column_letter(bcol)}{len(cl_rows) + 1}")
    for c, wdt in zip("ABCDEFGH", (8, 34, 14, 14, 60, 60, 26, 40)):
        ws.column_dimensions[c].width = wdt
    wp = wb.create_sheet("Personer")
    pcols = list(persons_out[0].keys())
    wp.append(pcols)
    for r in persons_out:
        wp.append([r[c] for c in pcols])
    wb.save(OUT_XLSX)
    print(f"\nWrote {os.path.relpath(OUT_XLSX, ROOT)}")


if __name__ == "__main__":
    main()
