"""place_v097_crosswalk.py — map V0.97 place records (LOC…) onto the V0.92
place IDs (L{LocID}) so V0.97 labels/coords/categories can overlay the
existing pipeline without re-keying anything.

Match order: exact name → case/diacritic-folded name → fuzzy (difflib >=0.88,
unique best). Unmatched rows on either side are listed for manual review.

Output: data/curated/place_v097_crosswalk.csv
Usage:  python scripts/normalization/place_v097_crosswalk.py
"""
import csv, difflib, re, unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
NEW = ROOT / "data/raw/4-LOCATION-V0.97.xlsx"
OLD = ROOT / "data/raw/HCA REPOSITORY V0.92/LocationData-PQ-V0.92.xlsx"
OUT = ROOT / "data/curated/place_v097_crosswalk.csv"
FIELDS = ["loc_id_v097", "label_v097", "entity_id_v092", "label_v092", "match", "score"]


def fold(t):
    t = unicodedata.normalize("NFKD", t.lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def load():
    new = [r for r in load_workbook(NEW, read_only=True)["Location"].iter_rows(min_row=5, values_only=True) if r[0]]
    old = [r for r in load_workbook(OLD, read_only=True)["DimLoc"].iter_rows(min_row=2, values_only=True) if r[0]]
    return new, old


def build():
    new, old = load()
    oid = lambda r: f"L{int(r[0]):05d}"
    by_exact = {r[1]: r for r in old}
    by_fold = {}
    for r in old:
        by_fold.setdefault(fold(r[1]), []).append(r)
    used, rows, pending = set(), [], []
    for n in new:
        o = by_exact.get(n[1])
        if o and oid(o) not in used:
            used.add(oid(o)); rows.append((n[0], n[1], oid(o), o[1], "exact", 1.0)); continue
        c = [x for x in by_fold.get(fold(n[1]), []) if oid(x) not in used]
        if len(c) == 1:
            used.add(oid(c[0])); rows.append((n[0], n[1], oid(c[0]), c[0][1], "folded", 1.0)); continue
        pending.append(n)
    left = {fold(r[1]): r for r in old if oid(r) not in used}
    for n in pending:
        m = difflib.get_close_matches(fold(n[1]), list(left), n=2, cutoff=0.88)
        if m and (len(m) == 1 or difflib.SequenceMatcher(None, fold(n[1]), m[0]).ratio()
                  - difflib.SequenceMatcher(None, fold(n[1]), m[1]).ratio() > 0.03):
            o = left.pop(m[0]); used.add(oid(o))
            rows.append((n[0], n[1], oid(o), o[1], "fuzzy", round(difflib.SequenceMatcher(None, fold(n[1]), m[0]).ratio(), 3)))
        else:
            rows.append((n[0], n[1], "", "", "new", 0))
    for o in left.values():
        rows.append(("", "", oid(o), o[1], "old-only", 0))
    return rows


def main():
    rows = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(FIELDS); w.writerows(rows)
    from collections import Counter
    print(Counter(r[4] for r in rows), "->", OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
