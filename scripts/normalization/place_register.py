"""place_register.py — the place register (STED-REGISTER) base file.

Since 2026-10-21 the base of the place register is the curated workbook

    data/HCA_Steder_GÆLDENDE_<DDMMYYYY>_LocationID_alle_<n>_steder.xlsx

(sheet "Ark1", one row per place, keyed by LocationID `LOC…`). It replaces
the V0.82 `Registry` place rows (`Reg…` ids), the V0.97 location workbook
and `Steder_i_dagboegerne_verificeret_udfyldt VER 1.0.xlsx` as the source
for place identity, label, category, country, coordinates and
visited/mentioned status. See docs/data-model/place-register-gaeldende.md.

Columns (header names are matched, not positions; a trailing `*` is
ignored):  Løbenummer, LocationID*, RegistryTitle, Category (Danish),
Country, Book, PageRef (= FIRST diary page), Latitude, Longitude,
GeocodingStatus (Fundet / Ikke fundet), Besøgt eller omtalt.

The workbook only carries the first diary page per place. The full list of
diary pages is the `FKDiaryPageID` column of data/raw/4-LOCATION-V0.97.xlsx
— the same LocationIDs (100 % overlap, verified 2026-10-26) and, for 2,428
of 2,433 places, a page list whose earliest page equals Book/PageRef. For
the five places where the new workbook corrected that first page the
corrected page replaces the V0.97 earliest page (see `merge_pages`).

Used by hca_xlsx_to_csv.py (entities/references/places.csv),
hca_v092_to_csv.py and place_v097_crosswalk.py. Stdlib + openpyxl only.
"""

import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

ROOT = Path(__file__).resolve().parents[2]
SEARCH_DIRS = [ROOT / "data", ROOT / "data" / "raw"]
# The Æ is a wildcard so the name matches however the OS normalises it.
PLACE_GLOB = "HCA_Steder_G*LDENDE_*.xlsx"
V097_XLSX = ROOT / "data" / "raw" / "4-LOCATION-V0.97.xlsx"

ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
ROMAN_TO_INT = {r: i for i, r in enumerate(ROMAN) if r}

# Danish `Category` in the workbook -> the English vocabulary that
# places-extra.js / places.html / kort.html have used since the verified
# category pass (docs/data-model/steder-verificeret-category-mapping.md).
# The English values stay the raw data; the UI translates them for display.
CATEGORY_DA_TO_EN = {
    "By": "City",
    "Ejendom": "Property",
    "Område": "Region",
    "Point of interest": "Point of interest (POI)",
    "Flod": "River",
    "Bjerg": "Mountain",
    "Ø": "Island",
    "Sø": "Lake",
    "Land": "Country",
    "Hav": "Sea",
    "Kirke": "Church",
    "Kontinent": "Continent",
    "Fiktivt sted": "Fictional place",   # new in the 2026-10-21 workbook
}

PLACE_FIELDS = ["entity_id", "lbnr", "label", "category_da", "category", "country",
                "book", "page", "lat", "lon", "geocoding_status", "visit_status"]


def _date_key(path: Path):
    """`..._21102026_...` (DDMMYYYY) -> (YYYY, MM, DD); (0,0,0) if absent."""
    m = re.search(r"_(\d{2})(\d{2})(\d{4})_", path.name)
    return (int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else (0, 0, 0)


def find_place_xlsx(explicit=None, required=True):
    """The newest `HCA_Steder_GÆLDENDE_*` workbook in data/ (or data/raw/).
    With required=False returns None instead of exiting when none exists."""
    if explicit:
        p = Path(explicit)
        if not p.exists():
            sys.exit(f"Place workbook not found: {p}")
        return p
    found = [p for d in SEARCH_DIRS if d.is_dir() for p in d.glob(PLACE_GLOB)]
    if not found:
        if not required:
            return None
        sys.exit(f"No place workbook matching {PLACE_GLOB} in "
                 f"{', '.join(str(d.relative_to(ROOT)) for d in SEARCH_DIRS)}")
    return max(found, key=lambda p: (_date_key(p), p.name))


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_places(path=None) -> list[dict]:
    """Place rows in register order = ascending LocationID. The ids are
    allocated in alphabetical slots (LOC1000000 Abbotsford, LOC1000050
    Abydos, …), so this is the printed-index order the Reg… ids had too; the
    workbook's own row order / `Løbenummer` (`lbnr`) is first-appearance order
    in the diaries and is kept only as a column. Rows without a LocationID
    (trailing blank rows) are skipped; duplicate ids abort the build."""
    path = find_place_xlsx(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [_s(c).rstrip("*").strip() for c in next(it)]
        col = {name: i for i, name in enumerate(header)}
        need = ["LocationID", "RegistryTitle", "Category", "Country", "Book",
                "PageRef", "Latitude", "Longitude", "GeocodingStatus"]
        missing = [n for n in need if n not in col]
        if missing:
            sys.exit(f"{path.name}: missing column(s) {missing}; found {header}")
        visit_col = next((i for n, i in col.items() if n.startswith("Besøgt")), None)
        out, seen = [], set()
        for row in it:
            loc = _s(row[col["LocationID"]])
            if not loc:
                continue
            if not re.fullmatch(r"LOC\d+", loc):
                sys.exit(f"{path.name}: malformed LocationID {loc!r}")
            if loc in seen:
                sys.exit(f"{path.name}: duplicate LocationID {loc}")
            seen.add(loc)
            lat, lon = _num(row[col["Latitude"]]), _num(row[col["Longitude"]])
            if lat is None or lon is None or (lat == 0 and lon == 0):
                lat = lon = None
            cat_da = _s(row[col["Category"]])
            book, page = row[col["Book"]], row[col["PageRef"]]
            out.append({
                "entity_id": loc,
                "lbnr": _s(row[col["Løbenummer"]]) if "Løbenummer" in col else str(len(out) + 1),
                "label": _s(row[col["RegistryTitle"]]),
                "category_da": cat_da,
                "category": CATEGORY_DA_TO_EN.get(cat_da, cat_da),
                "country": _s(row[col["Country"]]),
                "book": int(book) if _num(book) is not None else None,
                "page": int(page) if _num(page) is not None else None,
                "lat": lat, "lon": lon,
                "geocoding_status": _s(row[col["GeocodingStatus"]]),
                "visit_status": _s(row[visit_col]) if visit_col is not None else "",
            })
    finally:
        wb.close()
    out.sort(key=lambda p: int(p["entity_id"][3:]))
    return out


def unmapped_categories(places) -> set[str]:
    """Danish category values with no entry in CATEGORY_DA_TO_EN."""
    return {p["category_da"] for p in places
            if p["category_da"] and p["category_da"] not in CATEGORY_DA_TO_EN}


# ── diary page references ────────────────────────────────────────────────────

def _page_key(vp):
    return (ROMAN_TO_INT.get(vp[0], 99), int(vp[1]) if str(vp[1]).isdigit() else 0)


def load_v097_pages() -> dict[str, list[tuple[str, str]]]:
    """{LocationID: [(vol_roman, page), …]} from V0.97's FKDiaryPageID
    ("III-242 III-261 …"). Empty dict if the V0.97 workbook is absent."""
    if not V097_XLSX.exists():
        return {}
    wb = load_workbook(V097_XLSX, read_only=True, data_only=True)
    out = {}
    try:
        for row in wb["Location"].iter_rows(min_row=5, values_only=True):
            if not row[0] or not row[9]:
                continue
            pages = []
            for tok in str(row[9]).split():
                vol, _, pg = tok.partition("-")
                if vol in ROMAN_TO_INT and pg.isdigit():
                    pages.append((vol, pg))
            out[str(row[0]).strip()] = pages
    finally:
        wb.close()
    return out


def merge_pages(place, v097_pages):
    """Full page list for one place, in volume/page order. Book/PageRef of
    the current workbook is authoritative for the FIRST page: when it is
    not among the V0.97 pages it is a correction of the V0.97 earliest
    page, which it then replaces (Dornburg, Høgebjerg, Cuxhaven,
    Glücksborg, Lipariske Øer). Returns (pages, corrected: bool)."""
    pages = sorted(set(v097_pages), key=_page_key)
    if not place["book"] or not place["page"]:
        return pages, False
    first = (ROMAN[place["book"]], str(place["page"]))
    if first in pages:
        return pages, False
    corrected = bool(pages)
    if pages:
        pages = pages[1:]
    return sorted({first, *pages}, key=_page_key), corrected


def build_references(places, v097_pages=None):
    """references.csv rows for every place (V0.82 column layout).

    `page_id` is a per-row key (as in V0.82: PKRegistryPageID) — `PagL` +
    running number; `seq` is the running order in the register (LocationID =
    alphabetical order, then volume/page), the same meaning as V0.82's
    BookSeqNo."""
    v097_pages = load_v097_pages() if v097_pages is None else v097_pages
    rows, corrected = [], []
    for p in places:
        pages, was_corrected = merge_pages(p, v097_pages.get(p["entity_id"], []))
        if was_corrected:
            corrected.append(p)
        for vol, pg in pages:
            seq = len(rows) + 1
            rows.append({"page_id": f"PagL{seq:06d}", "entity_id": p["entity_id"],
                         "entity_label": p["label"], "vol": vol, "page": pg,
                         "seq": str(seq)})
    return rows, corrected


# ── entities.csv integration ────────────────────────────────────────────────

def place_entity_row(p) -> dict:
    """entities.csv row. `description` stays EMPTY on purpose: place.html
    prints it as a prose paragraph, and country/category travel in
    places.csv → places-extra.js instead."""
    return {"entity_id": p["entity_id"], "entity_type": "place",
            "category_h1": "STED-REGISTER", "genre_h2": "", "form_h3": "",
            "subform_h4": "", "label": p["label"], "description": "",
            "see": "", "see_also": "", "year_derived": "", "date_derived": "",
            "person_derived": ""}


def apply_place_base(entities, references, places=None, v097_pages=None):
    """Swap the place rows of (entities, references) for the place-register
    base. Non-place rows keep their order; the new place rows go where the
    first old place row was. Returns (entities, references, places,
    corrected_places, dropped_ref_count)."""
    places = load_places() if places is None else places
    new_refs, corrected = build_references(places, v097_pages)
    old_place_ids = {e["entity_id"] for e in entities if e["entity_type"] == "place"}
    out, inserted = [], False
    for e in entities:
        if e["entity_type"] == "place":
            if not inserted:
                out.extend(place_entity_row(p) for p in places)
                inserted = True
            continue
        out.append(e)
    if not inserted:
        out.extend(place_entity_row(p) for p in places)
    kept_refs = [r for r in references if r["entity_id"] not in old_place_ids]
    dropped = len(references) - len(kept_refs)
    return out, kept_refs + new_refs, places, corrected, dropped


def places_csv_row(p) -> dict:
    return {k: ("" if p[k] is None else p[k]) for k in PLACE_FIELDS}
