#!/usr/bin/env python3
"""Hand editor-approved contributions to the authoritative data
(docs/crowdsourcing.md §6).

Reads the submissions an editor has approved in the review interface and
writes the curated, tracked layer:

  data/curated/person_crowd_approved.tsv    machine-readable source (one row
                                            per accepted field; the shape of
                                            person_emendations.tsv plus
                                            provenance columns)
  data/curated/person_crowd_approved.xlsx   review surface for the Excel
                                            workflow: one row per person in the
                                            column layout of the 6-PERSON
                                            workbook, colour-coded by
                                            confidence, sources as cell
                                            comments (same pairing as
                                            personregister_xi_review_full.tsv/.xlsx)

and marks the submissions as incorporated. Nothing is ever written back into
the register's own columns: the approved data is a *layer on top* of the
printed register, with its source, confidence and contributor on every row
(the principle of docs/data-model/person-editorial-emendations.md).

    python scripts/crowdsource/export_approved.py            # report only
    python scripts/crowdsource/export_approved.py --write    # write + mark
    python scripts/crowdsource/export_approved.py --import-json fil.json --write
                                    # first take in a file from the browser-local mode

Re-running is safe: rows are keyed on (submission_id, field) and replaced.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import re
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import taxonomy as tx  # noqa: E402
from register import Register, ROOT, surname_of  # noqa: E402
from store import FileStore  # noqa: E402

CURATED = ROOT / "data" / "curated"
TSV_PATH = CURATED / "person_crowd_approved.tsv"
XLSX_PATH = CURATED / "person_crowd_approved.xlsx"
WORKBOOK = ROOT / "data" / "raw" / "6-PERSON-V0.97(5).xlsx"

COLUMNS = ["person_id", "match_surname", "match_refs", "field", "emend_field", "original", "emended",
           "source", "source_detail", "date", "confidence", "notes", "contributor", "contributor_id",
           "submission_id", "reviewer", "identification", "workbook_person_id", "guard"]
SOURCE_CODE = "CROWD"       # new row in the source-code table of person-editorial-emendations.md


def _norm_label(label: str) -> str:
    import unicodedata
    s = "".join(c for c in unicodedata.normalize("NFKD", label or "") if not unicodedata.combining(c))
    s = re.sub(r"\([^)]*\)", "", s.lower())
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", s)).strip()


def workbook_crosswalk(register: Register, path: Path = WORKBOOK) -> dict:
    """{Reg id → PER id of the 6-PERSON workbook} where the register label
    identifies exactly one row on each side. Anything ambiguous is left out
    rather than guessed — the editor then matches that row by hand."""
    try:
        import openpyxl
    except ImportError:
        return {}
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    by_label = defaultdict(list)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 4 or not row[0]:
            continue
        by_label[_norm_label(str(row[1] or ""))].append(str(row[0]))
    register._load()
    ours = defaultdict(list)
    for rid, (label, _d) in register._persons.items():
        ours[_norm_label(label)].append(rid)
    out = {}
    for key, rids in ours.items():
        if len(rids) == 1 and len(by_label.get(key, ())) == 1:
            out[rids[0]] = by_label[key][0]
    return out


def _original(field: str, snap: dict) -> str:
    label = snap.get("label", "")
    if field == "surname":
        return snap.get("surname") or surname_of(label)
    if field == "description":
        return snap.get("description", "")
    if field in ("birth_date", "death_date"):
        m = re.search(r"\((?:ca\.?\s*)?(\d{3,4})?\s*[–\-]\s*(\d{3,4})?\)", label)
        if m:
            return (m.group(1) if field == "birth_date" else m.group(2)) or ""
    return ""


def _source_detail(sub: dict, indices) -> str:
    parts = []
    for i in indices:
        s = sub["sources"][i]
        bit = s["citation"]
        if s.get("locator"):
            bit += f" [{s['locator']}]"
        if s.get("url"):
            bit += f" <{s['url']}>"
        if s.get("accessed"):
            bit += f" (opslået {s['accessed']})"
        if s.get("attachment"):
            bit += f" (bilag: {s['attachment']['file']})"
        parts.append(bit)
    return " | ".join(parts)


def build_rows(store: FileStore, register: Register, *, include_incorporated=False, crosswalk=None) -> list:
    """One row per accepted field (and per accepted identification statement)
    of every approved or partially approved submission."""
    rows = []
    for sub in sorted(store.all(), key=lambda s: s["id"]):
        if sub["status"] not in ("approved", "partially_approved"):
            continue
        if sub.get("incorporated") and not include_incorporated:
            continue
        rev = sub["review"]
        snap = sub.get("register_snapshot") or register.snapshot(sub["person_id"])
        now_snap = register.snapshot(sub["person_id"])
        guard = "ok"
        if now_snap and snap.get("label") and now_snap.get("label") != snap.get("label"):
            guard = f"register_changed: {snap['label']!r} -> {now_snap['label']!r}"
        base = {
            "person_id": sub["person_id"],
            "match_surname": now_snap.get("surname") or snap.get("surname", ""),
            "match_refs": now_snap.get("refs_signature") or snap.get("refs_signature", ""),
            "source": SOURCE_CODE,
            "date": (rev["decided"] or "")[:10],
            "contributor": sub["contributor"]["name"],
            "contributor_id": sub["contributor"]["id"],
            "submission_id": sub["id"],
            "reviewer": rev["reviewer"],
            "identification": sub["identification"]["status"],
            "workbook_person_id": (crosswalk or {}).get(sub["person_id"], ""),
            "guard": guard,
        }
        for i, f in enumerate(sub["fields"]):
            d = rev["fields"].get(str(i))
            if not d or d["decision"] not in tx.ACCEPTED_DECISIONS:
                continue
            spec = tx.FIELD_BY_KEY[f["field"]]
            notes = []
            if f["comment"]:
                notes.append("Bidragyder: " + f["comment"])
            if d.get("note"):
                notes.append("Redaktion: " + d["note"])
            if d["decision"] == "accept_edited":
                notes.append(f"Redaktionen har rettet værdien (bidragyderen skrev: {f['value']})")
            rows.append({**base, "field": f["field"], "emend_field": spec[3] or "",
                         "original": _original(f["field"], snap),
                         "emended": d.get("final_value") or f["value"],
                         "source_detail": _source_detail(sub, f["sources"]),
                         "confidence": d.get("final_confidence") or f["confidence"],
                         "notes": " ¶ ".join(notes)})
        ident = sub["identification"]
        if rev.get("identification") and rev["identification"]["decision"] == "accept":
            val = {"duplicate": f"duplicate:{ident['same_as']}"}.get(ident["status"], ident["status"])
            rows.append({**base, "field": "identification", "emend_field": "", "original": "",
                         "emended": val,
                         "source_detail": _source_detail(sub, range(len(sub["sources"]))),
                         "confidence": sub.get("overall_confidence") or "proposed",
                         "notes": " ¶ ".join(x for x in (ident["summary"], rev["identification"].get("note")) if x)})
    return rows


def _rel(p: Path) -> str:
    try:
        return str(p.relative_to(ROOT))
    except ValueError:
        return str(p)


def _cell(v) -> str:
    return str(v if v is not None else "").replace("\t", " ").replace("\r", " ").replace("\n", " ¶ ")


def rows_to_tsv(rows: list) -> str:
    out = io.StringIO()
    w = csv.writer(out, delimiter="\t", lineterminator="\n")
    w.writerow(COLUMNS)
    for r in rows:
        w.writerow([_cell(r.get(c)) for c in COLUMNS])
    return out.getvalue()


def read_tsv(path: Path | None = None) -> list:
    path = path or TSV_PATH
    if not path.exists():
        return []
    with open(path, encoding="utf-8", newline="") as f:
        return [{c: r.get(c, "") for c in COLUMNS} for r in csv.DictReader(f, delimiter="\t")]


def merge_rows(existing: list, new: list) -> list:
    keyed = {(r["submission_id"], r["field"]): r for r in existing}
    for r in new:
        keyed[(r["submission_id"], r["field"])] = r
    return sorted(keyed.values(), key=lambda r: (r["person_id"], r["submission_id"], r["field"]))


# ── Excel review surface ───────────────────────────────────────────────────

PERSON_SHEET = [
    # header,                         source field
    ("PersonID*",                     None),
    ("Register-ID",                   None),
    ("Person* (som i registret)",     None),
    ("Description (som i registret)", None),
    ("GivenNames",                    "given_names"),
    ("Surname",                       "surname"),
    ("NameVariants",                  "name_variants"),
    ("BirthDate",                     "birth_date"),
    ("BirthPlace",                    "birth_place"),
    ("DeathDate",                     "death_date"),
    ("DeathPlace",                    "death_place"),
    ("BirthYear",                     None),
    ("DeathYear",                     None),
    ("Gender",                        "gender"),
    ("Nationality",                   "nationality"),
    ("Residence",                     "residence"),
    ("Occupation",                    "occupation"),
    ("Role/Occupation",               "role_bucket"),
    ("Relations",                     "relations"),
    ("Description (supplement)",      "description"),
    ("Wikidata",                      "wikidata"),
    ("Authority (other)",             "authority_other"),
    ("Identification",                None),
    ("Weakest confidence",            None),
    ("Contributors",                  None),
    ("Submissions",                   None),
    ("Reviewer",                      None),
    ("Reviewed",                      None),
    ("Conflicts",                     None),
]
FILL = {"certain": "C6EFCE", "probable": "FFEB9C", "proposed": "F8CBAD"}
RANK = {"certain": 3, "probable": 2, "proposed": 1}


def rows_to_xlsx_bytes(rows: list, register: Register, crosswalk=None):
    try:
        import openpyxl
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None
    if crosswalk is None:
        crosswalk = workbook_crosswalk(register)
    wb = openpyxl.Workbook()

    # Sheet 1 — how to read this workbook
    ws0 = wb.active
    ws0.title = "Læsevejledning"
    for line in [
        "Godkendte bidrag fra crowdsourcing-prototypen (docs/crowdsourcing.md).",
        "",
        "Fanen »Personer« har én række pr. person i kolonneopbygningen fra 6-PERSON-V0.97 (PersonID*, Person*, Description, BirthYear, DeathYear, Nationality, Gender, Role/Occupation) med tilføjede kolonner til de nye oplysninger.",
        "Cellefarve = sikkerhedsgrad: grøn sikker · gul sandsynlig · orange forslag. Kilder, bidragyderens kommentar og redaktionens note står som celle-kommentar.",
        "Kun felter, redaktionen har godkendt, er med. Registrets egne kolonner er IKKE ændret: det er et lag ovenpå.",
        "PersonID* er kun udfyldt, hvor registerlabelen peger entydigt på én række i 6-PERSON; ellers kobles rækken manuelt via Register-ID.",
        "Fanen »Godkendte felter« er den samme tabel som data/curated/person_crowd_approved.tsv (kilden); denne fane er gennemsynsfladen.",
        "Flere godkendte værdier til samme felt vises i »Conflicts« — afgør dem før indlæsning.",
    ]:
        ws0.append([line])
    ws0.column_dimensions["A"].width = 140
    ws0["A1"].font = Font(bold=True, size=13)

    # Sheet 2 — one row per person
    ws = wb.create_sheet("Personer")
    ws.append([h for h, _ in PERSON_SHEET])
    by_person = defaultdict(list)
    for r in rows:
        by_person[r["person_id"]].append(r)
    for pid in sorted(by_person):
        claims = by_person[pid]
        label, desc = register.person(pid) or ("", "")
        values = {}
        conflicts = []
        for c in sorted(claims, key=lambda c: (RANK.get(c["confidence"], 0), c["date"]), reverse=True):
            if c["field"] in values:
                if values[c["field"]]["emended"].strip().lower() != c["emended"].strip().lower():
                    conflicts.append(f"{c['field']}: {c['emended']} ({c['contributor']})")
                continue
            values[c["field"]] = c
        ident = next((c for c in claims if c["field"] == "identification"), None)
        by_year = lambda k: re.sub(r"\D*(\d{4}).*", r"\1", values[k]["emended"]) if k in values else ""
        weakest = min((c["confidence"] for c in claims), key=lambda x: RANK.get(x, 0))
        out = []
        for header, key in PERSON_SHEET:
            if key:
                out.append(values[key]["emended"] if key in values else "")
            else:
                out.append({
                    "PersonID*": crosswalk.get(pid) or claims[0]["workbook_person_id"], "Register-ID": pid,
                    "Person* (som i registret)": label, "Description (som i registret)": desc,
                    "BirthYear": by_year("birth_date"), "DeathYear": by_year("death_date"),
                    "Identification": ident["emended"] if ident else claims[0]["identification"],
                    "Weakest confidence": weakest,
                    "Contributors": "; ".join(sorted({c["contributor"] for c in claims})),
                    "Submissions": "; ".join(sorted({c["submission_id"] for c in claims})),
                    "Reviewer": "; ".join(sorted({c["reviewer"] for c in claims})),
                    "Reviewed": max(c["date"] for c in claims),
                    "Conflicts": " | ".join(conflicts),
                }[header])
        ws.append(out)
        r_i = ws.max_row
        for col, (header, key) in enumerate(PERSON_SHEET, start=1):
            if key and key in values:
                c = values[key]
                cell = ws.cell(row=r_i, column=col)
                cell.fill = PatternFill("solid", fgColor=FILL.get(c["confidence"], "FFFFFF"))
                text = (f"{c['confidence']} · {c['contributor']} · {c['submission_id']}\n"
                        f"Kilder: {c['source_detail']}\n{c['notes']}")
                cell.comment = Comment(text[:1500], "crowd")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (h, _) in enumerate(PERSON_SHEET, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 22 if i > 4 else 30
        ws.cell(row=1, column=i).font = Font(bold=True)
    ws.freeze_panes = "C2"

    # Sheet 3 — the TSV
    ws3 = wb.create_sheet("Godkendte felter")
    ws3.append(COLUMNS)
    for r in rows:
        ws3.append([_cell(r.get(c)) for c in COLUMNS])
    for i in range(1, len(COLUMNS) + 1):
        ws3.cell(row=1, column=i).font = Font(bold=True)
    ws3.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--data-dir", default=str(ROOT / "data" / "crowdsource"))
    ap.add_argument("--write", action="store_true", help="write the curated files and mark submissions incorporated")
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--import-json", metavar="FILE", action="append", default=[],
                    help="first take in a contributions/decisions file made in the browser-local mode "
                         "(crowd.html → Mine bidrag, or crowd-review.html → Hent alle afgørelser)")
    args = ap.parse_args()

    store = FileStore(args.data_dir)
    register = Register()
    for path in args.import_json:
        import json
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        subs = data if isinstance(data, list) else data.get("submissions", [])
        counts = defaultdict(int)
        for obj in subs:
            try:
                counts[store.import_submission(obj, register.snapshot(obj.get("person_id", "")))] += 1
            except ValueError as e:
                counts["rejected"] += 1
                print(f"  [!] {e}")
        print(f"{path}: {dict(counts)}")
    crosswalk = workbook_crosswalk(register) if not args.no_xlsx else {}
    pending = build_rows(store, register, crosswalk=crosswalk)
    subs = sorted({r["submission_id"] for r in pending})
    print(f"{len(pending)} accepted field(s) from {len(subs)} approved submission(s) not yet incorporated")
    bad = [r for r in pending if r["guard"] != "ok"]
    for r in bad:
        print(f"  [!] {r['submission_id']} {r['person_id']}: {r['guard']}")
    if not args.write:
        print("(dry run — use --write to update data/curated/ and mark them incorporated)")
        return
    if bad:
        sys.exit("Refusing to write: the register changed under some contributions. "
                 "Re-review those submissions first (docs/crowdsourcing.md §6).")
    merged = merge_rows(read_tsv(), pending)
    CURATED.mkdir(parents=True, exist_ok=True)
    TSV_PATH.write_text(rows_to_tsv(merged), encoding="utf-8")
    print(f"wrote {_rel(TSV_PATH)}  ({len(merged)} rows)")
    if not args.no_xlsx:
        data = rows_to_xlsx_bytes(merged, register, crosswalk)
        if data is None:
            print("openpyxl not installed — .xlsx skipped")
        else:
            XLSX_PATH.write_bytes(data)
            print(f"wrote {_rel(XLSX_PATH)}")
    batch = "B-" + dt.date.today().isoformat()
    for sid in subs:
        store.mark_incorporated(sid, batch)
    print(f"marked {len(subs)} submission(s) incorporated ({batch}). "
          "Next: python scripts/build_all.py --only 4g   (rebuilds the site layer)")


if __name__ == "__main__":
    main()
